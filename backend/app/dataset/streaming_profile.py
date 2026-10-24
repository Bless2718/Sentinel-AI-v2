"""Bounded-memory file profiling for upload validation.

Duplicate counts are based on the first SAMPLE_LIMIT rows, to avoid
materializing all rows or keeping a full-dataset duplicate index in memory.
"""
import csv
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.dataset.profile import DatasetProfile

SAMPLE_LIMIT = 10000


def _profile_rows(columns: list[str], rows, source_bytes: int):
    width = len(columns)
    missing = 0
    row_count = 0
    seen: set[tuple[str, ...]] = set()
    sample_duplicates = 0
    kinds = [Counter() for _ in columns]

    for values in rows:
        row_count += 1
        values = list(values)
        if len(values) < width:
            values.extend([None] * (width - len(values)))
        values = values[:width]
        for i, value in enumerate(values):
            if value is None or (isinstance(value, str) and not value.strip()):
                missing += 1
            if row_count <= SAMPLE_LIMIT and value is not None:
                if isinstance(value, (date, datetime)):
                    kinds[i]["datetime"] += 1
                elif isinstance(value, (float, int)) and not isinstance(value, bool):
                    kinds[i]["number"] += 1
                else:
                    kinds[i]["string"] += 1
        if row_count <= SAMPLE_LIMIT:
            key = tuple(str(value) for value in values)
            if key in seen:
                sample_duplicates += 1
            else:
                seen.add(key)

    numeric = []
    datetimes = []
    categorical = []
    for name, kind in zip(columns, kinds):
        if kind["datetime"] and kind["datetime"] >= max(kind["number"], kind["string"]):
            datetimes.append(name)
        elif kind["number"] and kind["number"] >= kind["string"]:
            numeric.append(name)
        else:
            categorical.append(name)

    # Approximate duplicate rate from bounded sample, not an exact full-file count.
    sample_rows = min(row_count, SAMPLE_LIMIT)
    estimated_duplicates = round(sample_duplicates * row_count / sample_rows) if sample_rows else 0
    return DatasetProfile(
        rows=row_count,
        columns=width,
        missing_values=missing,
        duplicate_rows=estimated_duplicates,
        memory_usage_mb=round(source_bytes / (1024 * 1024), 2),
        numeric_columns=numeric,
        categorical_columns=categorical,
        datetime_columns=datetimes,
    )


def profile_file(path: Path) -> tuple[list[str], DatasetProfile]:
    size = path.stat().st_size
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            columns = [str(value).strip() if value is not None else "" for value in next(iterator, ())]
            if not columns:
                raise ValueError("Excel worksheet has no headers")
            return columns, _profile_rows(columns, iterator, size)
        finally:
            workbook.close()
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            columns = next(reader, [])
            if not columns:
                raise ValueError("CSV file has no headers")
            return columns, _profile_rows(columns, reader, size)
    raise ValueError("Unsupported dataset type")
